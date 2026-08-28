import json
from openpyxl import load_workbook, Workbook

RESULTS = json.load(open(r'D:\gdx\kaoshibao\debug\target_results.json', encoding='utf-8'))
SRC = r'D:\gdx\kaoshibao\output\考试宝题库_13573440_完整702.xlsx'
DST = r'D:\gdx\kaoshibao\output\考试宝题库_13573440_完整702.xlsx'

# 目标历史题号(1-based)
target_nums = [589, 639, 651, 653, 654]

wb = load_workbook(SRC)
ws = wb.active

# 找到这些题号对应的行并替换
replaced = 0
for row in ws.iter_rows(min_row=2):
    seq = row[0].value
    if seq in target_nums:
        key = str(seq)
        if key not in RESULTS:
            print(f"题{seq}: 未找到爬取结果")
            continue
        res = RESULTS[key]
        row[2].value = res['question']
        row[1].value = res['type'] or row[1].value
        row[9].value = res['answer'] or row[9].value
        row[10].value = res['analysis'] or row[10].value
        # 选项
        opts = res.get('options', [])
        for j in range(6):
            col_idx = 4 + j
            if col_idx <= 10:
                cell = row[col_idx - 1] if col_idx <= 10 else None
                if cell is not None:
                    cell.value = opts[j] if j < len(opts) else None
        replaced += 1
        print(f"题{seq} [{res['type']}] ans={res['answer']}: {str(res['question'])[:50]}")

wb.save(DST)
print(f"\n已替换 {replaced}/{len(target_nums)} 题 -> {DST}")

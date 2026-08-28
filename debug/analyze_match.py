import re
from openpyxl import load_workbook

def clean_title(t):
    t = re.sub(r'^(单选题|多选题|判断题|填空题|简答题)\s*\n?\d+/\d+、\s*\n?', '', t)
    t = re.sub(r'^\d+/\d+、\s*\n?', '', t)
    return t.strip()

def extract_alnum(t):
    return tuple(sorted(set(re.findall(r'[a-zA-Z0-9]{2,}', t))))

def extract_digits(t):
    return tuple(sorted(re.findall(r'\d+', t)))

cur_wb = load_workbook('output/考试宝题库_13573440_hybrid.xlsx', read_only=True)
cur_ws = cur_wb.active
cur = []
for i, row in enumerate(cur_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    t = clean_title(str(row[2]))
    cur.append({'alnum': extract_alnum(t), 'digits': extract_digits(t), 'len': len(t), 'ans': (str(row[9] or '')).strip(), 'type': str(row[1] or '')})
cur_wb.close()

hist_wb = load_workbook('output/考试宝题库_13573440_完整版-大数据IE笔试.xlsx', read_only=True)
hist_ws = hist_wb.active
hist = []
for i, row in enumerate(hist_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    t = clean_title(str(row[2]))
    hist.append({'alnum': extract_alnum(t), 'digits': extract_digits(t), 'len': len(t), 'ans': (str(row[4] or '')).strip()})
hist_wb.close()

no_alnum_no_digits = sum(1 for c in cur if not c['alnum'] and not c['digits'])
no_ans = sum(1 for c in cur if not c['ans'])
both = sum(1 for c in cur if not c['ans'] and not c['alnum'] and not c['digits'])
print(f'当前: alnum和digits都为空: {no_alnum_no_digits}, 无答案: {no_ans}, 都没有: {both}')

def match(c, h):
    if c['alnum'] and c['alnum'] == h['alnum']:
        return True
    if c['digits'] and c['digits'] == h['digits']:
        return True
    if c['ans'] and c['ans'] == h['ans'] and abs(c['len'] - h['len']) <= 2:
        return True
    return False

for i in range(min(20, len(cur), len(hist))):
    c, h = cur[i], hist[i]
    m = match(c, h)
    ca = 'Y' if c['alnum'] else 'N'
    cd = 'Y' if c['digits'] else 'N'
    ha = 'Y' if h['alnum'] else 'N'
    hd = 'Y' if h['digits'] else 'N'
    print(f'{i+1}: match={m}, cur(a={ca},d={cd},ans={c["ans"]},len={c["len"]}), hist(a={ha},d={hd},ans={h["ans"]},len={h["len"]})')

from openpyxl import load_workbook

cur_wb = load_workbook('output/考试宝题库_13573440_hybrid.xlsx', read_only=True)
cur_ws = cur_wb.active
cur = []
for i, row in enumerate(cur_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    cur.append((len(cur), str(row[1] or ''), str(row[2] or ''), (str(row[9] or '')).strip()))
cur_wb.close()

for kw in ['CDL', '特征预处理', '特征缩放', '离散化', 'NumPy', 'Pandas', 'Kafka', 'Broker', '结构化数据', '数据分析']:
    hits = [(i, t, a) for i, ty, t, a in cur if kw.lower() in t.lower()]
    print(f'--- "{kw}": {len(hits)} 条 ---')
    for i, t, a in hits[:8]:
        print(f'  [{i}] (ans={a}) {t[:60]}')

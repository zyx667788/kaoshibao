import re
from openpyxl import load_workbook, Workbook

def clean_title(t):
    t = re.sub(r'^(单选题|多选题|判断题|填空题|简答题)\s*\n?\d+/\d+、\s*\n?', '', t)
    t = re.sub(r'^\d+/\d+、\s*\n?', '', t)
    return t.strip()

def extract_alnum(t):
    return tuple(sorted(set(re.findall(r'[a-zA-Z0-9]{2,}', t))))

def extract_digits(t):
    return tuple(sorted(re.findall(r'\d+', t)))

def norm_ans(a):
    a = re.sub(r'[,，\s]+', '', a or '')
    return ''.join(sorted(set(c for c in a if c in 'ABCDEF')))

# 读当前698题
cur_wb = load_workbook('output/考试宝题库_13573440_hybrid.xlsx', read_only=True)
cur_ws = cur_wb.active
cur = []
for i, row in enumerate(cur_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    t = clean_title(str(row[2]))
    opts = []
    for j in range(4, 10):
        if j < len(row) and row[j]:
            opts.append(str(row[j]))
    ans = (str(row[9] or '')).strip() if len(row) > 9 else ''
    cur.append({
        '题目类型': str(row[1] or ''),
        '题目': str(row[2] or ''),
        '选项': opts,
        '正确答案': ans,
        '解析': str(row[10] or '') if len(row) > 10 else '',
        '_alnum': extract_alnum(t),
        '_digits': extract_digits(t),
        '_len': len(t),
        '_nans': norm_ans(ans),
    })
cur_wb.close()

# 读历史702题
hist_wb = load_workbook('output/考试宝题库_13573440_完整版-大数据IE笔试.xlsx', read_only=True)
hist_ws = hist_wb.active
hist = []
for i, row in enumerate(hist_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    title_raw = str(row[2])
    t = clean_title(title_raw)
    ans = (str(row[4] or '')).strip() if len(row) > 4 else ''
    hist.append({
        '题目类型': '',
        '题目': title_raw,
        '选项': str(row[3] or '') if len(row) > 3 else '',
        '正确答案': ans,
        '解析': str(row[5] or '') if len(row) > 5 else '',
        '_alnum': extract_alnum(t),
        '_digits': extract_digits(t),
        '_len': len(t),
        '_nans': norm_ans(ans),
    })
hist_wb.close()

print(f"当前: {len(cur)} 题, 有答案: {sum(1 for q in cur if q['_nans'])}")
print(f"历史: {len(hist)} 题, 有答案: {sum(1 for q in hist if q['_nans'])}")

def match(c, h):
    # alnum都非空且相同
    if c['_alnum'] and h['_alnum'] and c['_alnum'] == h['_alnum']:
        return True
    # digits都非空且相同
    if c['_digits'] and h['_digits'] and c['_digits'] == h['_digits']:
        return True
    # 标准化答案相同且长度差<=3
    if c['_nans'] and c['_nans'] == h['_nans'] and abs(c['_len'] - h['_len']) <= 3:
        return True
    return False

# DP: LCS
n = len(cur)
m = len(hist)
print(f"DP表: {n} x {m}")

dp = [[0] * (m + 1) for _ in range(n + 1)]
for i in range(1, n + 1):
    for j in range(1, m + 1):
        if match(cur[i-1], hist[j-1]):
            dp[i][j] = dp[i-1][j-1] + 1
        else:
            dp[i][j] = max(dp[i-1][j], dp[i][j-1])

print(f"LCS长度: {dp[n][m]}")

# 回溯
matches = []
i, j = n, m
while i > 0 and j > 0:
    if match(cur[i-1], hist[j-1]) and dp[i][j] == dp[i-1][j-1] + 1:
        matches.append((i-1, j-1))
        i -= 1
        j -= 1
    elif dp[i-1][j] >= dp[i][j-1]:
        i -= 1
    else:
        j -= 1
matches.reverse()

print(f"匹配对: {len(matches)}")
matched_hist = set(h for c, h in matches)
missing_hist = sorted(set(range(m)) - matched_hist)
print(f"历史未匹配: {len(missing_hist)} 题 (题号: {[h+1 for h in missing_hist]})")

# 构建最终702题：以历史顺序为基底
result = []
for hi in range(m):
    ci = None
    for c, h in matches:
        if h == hi:
            ci = c
            break
    if ci is not None:
        q = dict(cur[ci])
        q['_source'] = 'current'
        if not q.get('正确答案') and hist[hi].get('正确答案'):
            q['正确答案'] = hist[hi]['正确答案']
            q['解析'] = hist[hi].get('解析', '')
        result.append(q)
    else:
        q = dict(hist[hi])
        q['选项'] = [q['选项']] if q.get('选项') else []
        q['_source'] = 'history'
        result.append(q)

for i, q in enumerate(result):
    q['序号'] = i + 1

has_ans = sum(1 for q in result if q.get('正确答案'))
from_current = sum(1 for q in result if q.get('_source') == 'current')
from_hist = sum(1 for q in result if q.get('_source') == 'history')
print(f"最终: {len(result)} 题, 有答案: {has_ans}, 当前: {from_current}, 历史: {from_hist}")

# 保存
wb2 = Workbook()
ws2 = wb2.active
ws2.title = "题库"
headers = ['序号', '题目类型', '题目', '选项A', '选项B', '选项C', '选项D', '选项E', '选项F', '正确答案', '解析']
for col, h in enumerate(headers, 1):
    ws2.cell(row=1, column=col, value=h)
for q in result:
    r = q['序号'] + 1
    ws2.cell(row=r, column=1, value=q['序号'])
    ws2.cell(row=r, column=2, value=q.get('题目类型', ''))
    ws2.cell(row=r, column=3, value=q.get('题目', ''))
    opts = q.get('选项', [])
    if isinstance(opts, list):
        for j, opt in enumerate(opts):
            if j < 6:
                ws2.cell(row=r, column=4+j, value=opt)
    else:
        ws2.cell(row=r, column=4, value=opts)
    ws2.cell(row=r, column=10, value=q.get('正确答案', ''))
    ws2.cell(row=r, column=11, value=q.get('解析', ''))
wb2.save('output/考试宝题库_13573440_完整702.xlsx')
print("已保存: output/考试宝题库_13573440_完整702.xlsx")

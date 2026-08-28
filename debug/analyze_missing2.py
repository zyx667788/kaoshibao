import re
from openpyxl import load_workbook

def clean_title(t):
    t = re.sub(r'^(单选题|多选题|判断题|填空题|简答题)\s*\n?\d+/\d+、\s*\n?', '', t)
    t = re.sub(r'^\d+/\d+、\s*\n?', '', t)
    return t.strip()

def norm_ans(a):
    a = re.sub(r'[,，\s]+', '', a or '')
    return ''.join(sorted(set(c for c in a if c in 'ABCDEF')))

def extract_alnum(t):
    return tuple(sorted(set(re.findall(r'[a-zA-Z0-9]{2,}', t))))

def extract_digits(t):
    return tuple(sorted(re.findall(r'\d+', t)))

# 当前697题
cur_wb = load_workbook('output/考试宝题库_13573440_hybrid.xlsx', read_only=True)
cur_ws = cur_wb.active
cur = []
for i, row in enumerate(cur_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    t = clean_title(str(row[2]))
    ans = (str(row[9] or '')).strip() if len(row) > 9 else ''
    opts = []
    for j in range(4, 10):
        if j < len(row) and row[j]:
            opts.append(str(row[j]))
    cur.append({'nans': norm_ans(ans), 'len': len(t), 'alnum': extract_alnum(t), 'digits': extract_digits(t), 'title': str(row[2] or '')[:45], 'row': row})
cur_wb.close()

# 历史702题
hist_wb = load_workbook('output/考试宝题库_13573440_完整版-大数据IE笔试.xlsx', read_only=True)
hist_ws = hist_wb.active
hist = []
for i, row in enumerate(hist_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    t = clean_title(str(row[2]))
    ans = (str(row[4] or '')).strip() if len(row) > 4 else ''
    hist.append({'nans': norm_ans(ans), 'len': len(t), 'alnum': extract_alnum(t), 'digits': extract_digits(t), 'title': str(row[2] or '')[:45]})
hist_wb.close()

missing_hist = [1, 239, 300, 335, 588, 657, 659, 663, 664, 667, 675, 676, 678, 684, 686, 689, 692, 693, 694, 696, 697, 698]

print('=== 22题历史未匹配：在当前中找候选 ===')
for hi in missing_hist:
    h = hist[hi]
    candidates = []
    for ci, c in enumerate(cur):
        score = 0
        if c['nans'] and c['nans'] == h['nans']:
            score += 10
        if c['alnum'] and h['alnum'] and c['alnum'] == h['alnum']:
            score += 20
        if c['digits'] and h['digits'] and c['digits'] == h['digits']:
            score += 15
        if abs(c['len'] - h['len']) <= 3:
            score += 5
        if score >= 20:
            candidates.append((ci, score, c['len'], c['title']))
    candidates.sort(key=lambda x: -x[1])
    if candidates:
        print(f'历史{hi+1}: ans={h["nans"]}, len={h["len"]}, "{h["title"]}"')
        for cand in candidates[:3]:
            print(f'   候选[{cand[0]}] score={cand[1]} len={cand[2]} "{cand[3]}"')
    else:
        print(f'历史{hi+1}: ans={h["nans"]}, len={h["len"]}, "{h["title"]}" -> 无候选')

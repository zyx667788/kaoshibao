import re
from openpyxl import load_workbook

cur_wb = load_workbook('output/考试宝题库_13573440_hybrid.xlsx', read_only=True)
cur_ws = cur_wb.active
cur = []
for i, row in enumerate(cur_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    cur.append(str(row[2] or ''))
cur_wb.close()

keywords = ['点击上方设置', '频率就是概率', '原始数据', '批量梯度下降', '手肘法', '朴素贝叶斯', '特征选择', '数据治理', '关联分析', '大数据平台系统架构', '风险结构', '欠拟合', '训练数据较少', '典型降维算法', '逐层搜索']
for kw in keywords:
    hits = [t[:60] for t in cur if kw in t]
    print(f'"{kw}": {len(hits)} 条')
    for h in hits[:3]:
        print(f'   {h}')

import json, re
from openpyxl import load_workbook, Workbook

def clean_title(t):
    t = re.sub(r'^(单选题|多选题|判断题|填空题|简答题)\s*\n?\d+/\d+、\s*\n?', '', t)
    t = re.sub(r'^\d+/\d+、\s*\n?', '', t)
    return t.strip()

def extract_alnum(t):
    return tuple(sorted(set(re.findall(r'[a-zA-Z0-9]{2,}', t))))

def extract_numbers(t):
    return tuple(sorted(set(re.findall(r'\d+', t))))

wb = load_workbook('output/考试宝题库_13573440_完整版-大数据IE笔试.xlsx', read_only=True)
ws = wb.active
hist = []
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    title_raw = str(row[2])
    title_clean = clean_title(title_raw)
    ans = (str(row[4] or '')).strip() if len(row) > 4 else ''
    hist.append({
        '序号': len(hist) + 1,
        '题目类型': str(row[1] or ''),
        '题目': title_raw,
        '选项': str(row[3] or '') if len(row) > 3 else '',
        '正确答案': ans,
        '解析': str(row[5] or '') if len(row) > 5 else '',
        '_alnum': extract_alnum(title_clean),
        '_nums': extract_numbers(title_clean),
        '_len': len(title_clean),
        '_ans': ans,
        '_matched': False,
    })
wb.close()

current = json.load(open('data/progress_13573440.json', encoding='utf-8'))
for q in current:
    t = clean_title(q.get('题目', ''))
    q['_alnum'] = extract_alnum(t)
    q['_nums'] = extract_numbers(t)
    q['_len'] = len(t)
    q['_ans'] = (q.get('正确答案', '') or '').strip()

matched_count = 0
for cq in current:
    best = None
    best_score = 0
    for hq in hist:
        if hq['_matched']:
            continue
        if hq['_ans'] != cq['_ans']:
            continue
        score = 0
        if hq['_alnum'] == cq['_alnum'] and cq['_alnum']:
            score += 10
        elif set(hq['_alnum']) & set(cq['_alnum']):
            score += len(set(hq['_alnum']) & set(cq['_alnum']))
        if abs(hq['_len'] - cq['_len']) <= 3:
            score += 3
        elif abs(hq['_len'] - cq['_len']) <= 10:
            score += 1
        if score > best_score:
            best_score = score
            best = hq
    if best and best_score >= 3:
        best['_matched'] = True
        best['题目'] = cq.get('题目', '')
        best['题目类型'] = cq.get('题目类型', best['题目类型'])
        opts = cq.get('选项', [])
        best['选项'] = opts if isinstance(opts, list) else [opts]
        best['正确答案'] = cq.get('正确答案', best['正确答案'])
        best['解析'] = cq.get('解析', best['解析'])
        best['_source'] = 'current'
        matched_count += 1

for hq in hist:
    if not hq.get('_source'):
        hq['_source'] = 'history'

print(f"历史: {len(hist)} 题")
print(f"当前: {len(current)} 题")
print(f"匹配成功(用当前文字替换): {matched_count} 题")
print(f"未匹配(保留历史): {len(hist) - matched_count} 题")

final = sorted(hist, key=lambda q: q['序号'])
has_ans = sum(1 for q in final if q.get('正确答案'))
from_current = sum(1 for q in final if q.get('_source') == 'current')
print(f"最终: {len(final)} 题, 有答案: {has_ans}, 来自当前: {from_current}, 来自历史: {len(final)-from_current}")

wb2 = Workbook()
ws2 = wb2.active
ws2.title = "题库"
headers = ['序号', '题目类型', '题目', '选项A', '选项B', '选项C', '选项D', '选项E', '选项F', '正确答案', '解析']
for col, h in enumerate(headers, 1):
    ws2.cell(row=1, column=col, value=h)
for q in final:
    r = q['序号'] + 1
    ws2.cell(row=r, column=1, value=q['序号'])
    ws2.cell(row=r, column=2, value=q.get('题目类型', ''))
    ws2.cell(row=r, column=3, value=q.get('题目', ''))
    opts = q.get('选项', [])
    if isinstance(opts, list):
        for j, opt in enumerate(opts):
            if j < 6:
                ws2.cell(row=r, column=4+j, value=opt)
    else:
        ws2.cell(row=r, column=4, value=opts)
    ws2.cell(row=r, column=10, value=q.get('正确答案', ''))
    ws2.cell(row=r, column=11, value=q.get('解析', ''))
wb2.save('output/考试宝题库_13573440_完整702.xlsx')
print("已保存: output/考试宝题库_13573440_完整702.xlsx")

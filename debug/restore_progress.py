import json
from openpyxl import load_workbook

PAPER_ID = '13573440'
wb = load_workbook(r'D:\gdx\kaoshibao\output\考试宝题库_13573440_hybrid.xlsx', read_only=True)
ws = wb.active
questions = []
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    opts = []
    for j in range(4, 10):
        if j < len(row) and row[j]:
            opts.append(str(row[j]))
    questions.append({
        '序号': len(questions) + 1,
        '题目类型': str(row[1] or ''),
        '题目': str(row[2] or ''),
        '选项': opts,
        '正确答案': (str(row[9] or '')).strip() if len(row) > 9 else '',
        '解析': str(row[10] or '') if len(row) > 10 else '',
    })
wb.close()

with open(rf'D:\gdx\kaoshibao\data\progress_{PAPER_ID}_hybrid.json', 'w', encoding='utf-8') as f:
    json.dump(questions, f, ensure_ascii=False, indent=2)
print(f"已恢复进度: {len(questions)} 题 -> progress_{PAPER_ID}_hybrid.json")

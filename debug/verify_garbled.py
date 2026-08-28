from openpyxl import load_workbook

# 明确乱码词（来自历史DOM版防爬字体的典型替换）
GARBLED_WORDS = ['欢据', '少据', '丈章', '袋场', '抱法', '忍性', '士选代', '医', '玛Hudi', '玛CDL', '玛于',
                 '东民', '派台', '拥担', '谈那度', '场冷', '万习', '出叶', '父程', '石金融', '独构', '照务',
                 '时西', '交衣', '令用', '委成', '飞标', '于用', '园法', '且述', '健项', '首下', '暗下',
                 '伦哪', '文自', '兵数', '令士', '止关', '名觉', '忠领', '员案', '惜题', '医主', '委化']

wb = load_workbook('output/考试宝题库_13573440_完整702.xlsx', read_only=True)
ws = wb.active
total = 0
found = []
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    total += 1
    title = str(row[2] or '')
    hits = [w for w in GARBLED_WORDS if w in title]
    if hits:
        found.append((row[0], hits, title[:50]))
wb.close()

print(f'总题数: {total}')
print(f'含明确乱码词题数: {len(found)}')
for n, hits, tt in found[:20]:
    print(f'  题{n} {hits}: {tt}')

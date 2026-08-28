from openpyxl import load_workbook

wb = load_workbook('output/考试宝题库_13573440_完整702.xlsx', read_only=True)
ws = wb.active
total = has_ans = 0
types = {}
garbled = []
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    total += 1
    if str(row[9] or '').strip():
        has_ans += 1
    t = str(row[1] or '')
    types[t] = types.get(t, 0) + 1
    title = str(row[2] or '')
    # 检查常见乱码字
    bad = [c for c in ['士', '玛', '医', '永', '忍', '袋', '抱', '少', '从', '基', '族', '喜', '宴', '爸', '冫', '灬'] if c in title]
    if bad:
        garbled.append((row[0], bad, title[:40]))
wb.close()

print(f'总题数: {total}, 有答案: {has_ans}')
print(f'类型分布: {types}')
print(f'含乱码字题数: {len(garbled)}')
for n, bad, tt in garbled[:15]:
    print(f'  题{n} {bad}: {tt}')

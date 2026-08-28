import re
from difflib import SequenceMatcher
from openpyxl import load_workbook, Workbook

def clean_title(t):
    t = re.sub(r'^(单选题|多选题|判断题|填空题|简答题)\s*\n?\d+/\d+、\s*\n?', '', t)
    t = re.sub(r'^\d+/\d+、\s*\n?', '', t)
    return t.strip()

def norm_ans(a):
    a = re.sub(r'[,，\s]+', '', a or '')
    return ''.join(sorted(set(c for c in a if c in 'ABCDEF')))

def extract_alnum(t):
    return tuple(sorted(set(re.findall(r'[a-zA-Z0-9]{2,}', t))))

def extract_digits(t):
    return tuple(sorted(re.findall(r'\d+', t)))

def extract_words(t):
    return set(re.findall(r'[a-zA-Z]{3,}', t))

def extract_cn(t):
    return re.sub(r'[\s\dA-Za-z，。、：；！？（）()【】\[\]\'"“”‘’—\-—,.!?:;（）]+', '', t)

def compact(t):
    return re.sub(r'[\s，。、：；！？（）()【】\[\]\'"“”‘’—\-—,.!?:;]+', '', t)

def ph_count(t):
    return t.count('（ ）') + t.count('()') + t.count('（  ）')

# 读当前697题
cur_wb = load_workbook('output/考试宝题库_13573440_hybrid.xlsx', read_only=True)
cur_ws = cur_wb.active
cur = []
for i, row in enumerate(cur_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    t = clean_title(str(row[2]))
    opts = []
    for j in range(4, 10):
        if j < len(row) and row[j]:
            opts.append(str(row[j]))
    ans = (str(row[9] or '')).strip() if len(row) > 9 else ''
    cur.append({
        '题目类型': str(row[1] or ''),
        '题目': str(row[2] or ''),
        '选项': opts,
        '正确答案': ans,
        '解析': str(row[10] or '') if len(row) > 10 else '',
        '_alnum': extract_alnum(t),
        '_digits': extract_digits(t),
        '_words': extract_words(t),
        '_compact': compact(t),
        '_cn': extract_cn(t),
        '_ph': ph_count(t),
        '_len': len(t),
        '_nans': norm_ans(ans),
        '_used': False,
    })
cur_wb.close()

# 读历史702题
hist_wb = load_workbook('output/考试宝题库_13573440_完整版-大数据IE笔试.xlsx', read_only=True)
hist_ws = hist_wb.active
hist = []
for i, row in enumerate(hist_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    title_raw = str(row[2])
    t = clean_title(title_raw)
    ans = (str(row[4] or '')).strip() if len(row) > 4 else ''
    hist.append({
        '题目类型': '',
        '题目': title_raw,
        '选项': str(row[3] or '') if len(row) > 3 else '',
        '正确答案': ans,
        '解析': str(row[5] or '') if len(row) > 5 else '',
        '_alnum': extract_alnum(t),
        '_digits': extract_digits(t),
        '_words': extract_words(t),
        '_compact': compact(t),
        '_cn': extract_cn(t),
        '_ph': ph_count(t),
        '_len': len(t),
        '_nans': norm_ans(ans),
    })
hist_wb.close()

print(f"当前: {len(cur)} 题, 历史: {len(hist)} 题")

def match_score(c, h):
    """返回匹配分数，0=不匹配"""
    s = 0
    if c['_alnum'] and c['_alnum'] == h['_alnum']:
        s += 100
    if c['_alnum'] and h['_alnum']:
        inter = set(c['_alnum']) & set(h['_alnum'])
        if inter and len(inter) >= max(1, min(len(c['_alnum']), len(h['_alnum'])) - 1):
            s += 60
    if c['_words'] and h['_words']:
        inter = c['_words'] & h['_words']
        if inter:
            s += 30
    if c['_nans'] and c['_nans'] == h['_nans'] and abs(c['_len'] - h['_len']) <= 5:
        s += 40
    if c['_digits'] and c['_digits'] == h['_digits']:
        s += 30
    return s

# 第一轮：强匹配
used_cur = set()
hist_to_cur = {}

# 手工精确映射：历史题号(1-based) -> 当前索引(0-based)，优先占用
MANUAL = {
    621: 615, 632: 626, 645: 639, 655: 649, 656: 650, 666: 660, 670: 664,
    677: 671, 678: 672, 681: 675, 683: 677, 684: 678, 693: 687, 694: 688,
    695: 689, 696: 690, 697: 691, 698: 692, 699: 693,
}
for hnum, cidx in MANUAL.items():
    hist_to_cur[hnum - 1] = cidx
    used_cur.add(cidx)

for hi, h in enumerate(hist):
    if hi in hist_to_cur:
        continue
    best_score = 0
    best_ci = -1
    for ci, c in enumerate(cur):
        if ci in used_cur:
            continue
        s = match_score(c, h)
        if s > best_score:
            best_score = s
            best_ci = ci
    if best_score >= 40 and best_ci >= 0:
        hist_to_cur[hi] = best_ci
        used_cur.add(best_ci)

# 第二轮：弱匹配（填空/简答无答案题，用占位符数+英文token+中文相似度）
for hi, h in enumerate(hist):
    if hi in hist_to_cur:
        continue
    if h['_ph'] == 0 and not h['_alnum'] and not h['_words'] and not h['_digits']:
        continue
    best_score = 0
    best_ci = -1
    for ci, c in enumerate(cur):
        if ci in used_cur:
            continue
        s = 0
        if h['_ph'] and h['_ph'] == c['_ph']:
            s += 15
        if h['_words'] and h['_words'] & c['_words']:
            s += 25
        if h['_alnum'] and set(h['_alnum']) & set(c['_alnum']):
            s += 20
        if h['_digits'] and set(h['_digits']) & set(c['_digits']):
            s += 10
        ratio = SequenceMatcher(None, h['_compact'][:40], c['_compact'][:40]).ratio()
        if ratio >= 0.5:
            s += int(ratio * 50)
        if s > best_score:
            best_score = s
            best_ci = ci
    if best_score >= 25 and best_ci >= 0:
        hist_to_cur[hi] = best_ci
        used_cur.add(best_ci)

print(f"匹配成功: {len(hist_to_cur)}/{len(hist)}, 已用当前题: {len(used_cur)}/{len(cur)}")
missing = [hi for hi in range(len(hist)) if hi not in hist_to_cur]
print(f"历史未匹配: {len(missing)} 题, 题号: {[m+1 for m in missing]}")

# 构建最终702：以历史顺序为基底
result = []
for hi, h in enumerate(hist):
    if hi in hist_to_cur:
        ci = hist_to_cur[hi]
        c = cur[ci]
        q = dict(c)
        q['_source'] = 'current'
        if not q.get('正确答案') and h.get('正确答案'):
            q['正确答案'] = h['正确答案']
            q['解析'] = h.get('解析', '')
        result.append(q)
    else:
        q = dict(h)
        q['选项'] = [q['选项']] if q.get('选项') else []
        q['_source'] = 'history'
        mt = re.match(r'^(单选题|多选题|判断题|填空题|简答题)', str(q.get('题目', '')).strip())
        if mt:
            q['题目类型'] = mt.group(1)
        result.append(q)

for i, q in enumerate(result):
    q['序号'] = i + 1

has_ans = sum(1 for q in result if q.get('正确答案'))
from_current = sum(1 for q in result if q.get('_source', '').startswith('current'))
from_hist = sum(1 for q in result if q.get('_source') == 'history')
print(f"最终: {len(result)} 题, 有答案: {has_ans}, 当前文字: {from_current}, 历史文字: {from_hist}")

wb2 = Workbook()
ws2 = wb2.active
ws2.title = "题库"
headers = ['序号', '题目类型', '题目', '选项A', '选项B', '选项C', '选项D', '选项E', '选项F', '正确答案', '解析']
for col, h in enumerate(headers, 1):
    ws2.cell(row=1, column=col, value=h)
for q in result:
    r = q['序号'] + 1
    ws2.cell(row=r, column=1, value=q['序号'])
    ws2.cell(row=r, column=2, value=q.get('题目类型', ''))
    ws2.cell(row=r, column=3, value=q.get('题目', ''))
    opts = q.get('选项', [])
    if isinstance(opts, list):
        for j, opt in enumerate(opts):
            if j < 6:
                ws2.cell(row=r, column=4+j, value=opt)
    else:
        ws2.cell(row=r, column=4, value=opts)
    ws2.cell(row=r, column=10, value=q.get('正确答案', ''))
    ws2.cell(row=r, column=11, value=q.get('解析', ''))
wb2.save('output/考试宝题库_13573440_完整702.xlsx')
print("已保存: output/考试宝题库_13573440_完整702.xlsx")

# 打印历史文字补充的题
print()
print("=== 历史文字补充的题（需要检查乱码）===")
for q in result:
    if q.get('_source') == 'history':
        print(f"题{q['序号']}: {str(q['题目'])[:60]} | 答案: {q.get('正确答案', '')}")

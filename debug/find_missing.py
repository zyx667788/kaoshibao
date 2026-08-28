import re
from openpyxl import load_workbook

def clean_title(t):
    t = re.sub(r'^(单选题|多选题|判断题|填空题|简答题)\s*\n?\d+/\d+、\s*\n?', '', t)
    t = re.sub(r'^\d+/\d+、\s*\n?', '', t)
    return t.strip()

# 当前697题
cur_wb = load_workbook('output/考试宝题库_13573440_hybrid.xlsx', read_only=True)
cur_ws = cur_wb.active
cur_titles = []
for i, row in enumerate(cur_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    cur_titles.append(clean_title(str(row[2])))
cur_wb.close()

# 历史702题
hist_wb = load_workbook('output/考试宝题库_13573440_完整版-大数据IE笔试.xlsx', read_only=True)
hist_ws = hist_wb.active
hist = []
for i, row in enumerate(hist_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    t = clean_title(str(row[2]))
    ans = (str(row[4] or '')).strip() if len(row) > 4 else ''
    hist.append({'num': len(hist) + 1, 'title': t, 'ans': ans})
hist_wb.close()

def extract_keys(t):
    words = set(re.findall(r'[a-zA-Z]{3,}', t))
    nums = set(re.findall(r'\d+', t))
    ph = t.count('（ ）') + t.count('()')
    return words, nums, ph

def find_in_cur(t):
    """在当前中找包含关键词的题，返回命中数"""
    words, nums, ph = extract_keys(t)
    hits = 0
    for ct in cur_titles:
        score = 0
        if words:
            inter = words & set(re.findall(r'[a-zA-Z]{3,}', ct))
            if inter:
                score += len(inter)
        if nums:
            inter = nums & set(re.findall(r'\d+', ct))
            if inter:
                score += len(inter)
        if score >= 1:
            hits += 1
    return hits

# 对每个历史题，检查当前是否有对应（关键词命中）
print('=== 历史题在当前697中无关键词命中的（可能缺失）===')
no_hit = []
for h in hist:
    hits = find_in_cur(h['title'])
    if hits == 0:
        no_hit.append(h)
        print(f"历史{h['num']} (ans={h['ans']}): {h['title'][:55]}")

print(f"\n无命中总数: {len(no_hit)}")

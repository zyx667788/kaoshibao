from openpyxl import load_workbook

cur_wb = load_workbook('output/考试宝题库_13573440_hybrid.xlsx', read_only=True)
cur_ws = cur_wb.active
cur = []
for i, row in enumerate(cur_ws.iter_rows(values_only=True)):
    if i == 0 or not row[2]:
        continue
    cur.append((len(cur), str(row[1] or ''), str(row[2] or ''), (str(row[9] or '')).strip()))
cur_wb.close()

for kw in ['os.path', 'StandardScaler', 'PySpark', 'exists', 'Kafka的吞吐率', '关注样本方差', '物理上把Topic']:
    hits = [(i, t, a) for i, ty, t, a in cur if kw.lower() in t.lower()]
    print(f'--- "{kw}": {len(hits)} 条 ---')
    for i, t, a in hits[:6]:
        print(f'  [{i}] (ans={a}) {t[:70]}')
